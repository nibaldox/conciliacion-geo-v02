"""Tests for core.word_report_layout — presentation helper for the Word report."""
from datetime import datetime

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt, RGBColor

from core.word_report_layout import add_cover, configure_document, style_data_table


def _cell_fill(cell):
    tcPr = cell._tc.tcPr
    if tcPr is None:
        return None
    for child in tcPr:
        if child.tag == qn("w:shd"):
            return child.get(qn("w:fill"))
    return None


def _make_data_table(doc, rows=None):
    rows = rows or [
        ["Sección", "H. Real (m)", "Estado"],
        ["S-01", "15.2", "✓ 85"],
        ["S-02", "16.1", "✗ 40"],
        ["S-03", "N/A", "✓ 90"],
    ]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row_values in enumerate(rows):
        for c, value in enumerate(row_values):
            table.rows[r].cells[c].text = str(value)
    return table


def _row_has_flag(row, tag):
    trPr = row._tr.trPr
    return trPr is not None and trPr.find(qn(tag)) is not None


MM_TOLERANCE = Mm(0.5)


def _assert_mm(actual, expected_mm):
    assert abs(actual - Mm(expected_mm)) < MM_TOLERANCE


class TestConfigureDocumentPageSetup:
    def test_landscape_a4_with_uniform_margins(self):
        doc = Document()
        configure_document(doc)
        section = doc.sections[0]
        assert section.orientation == WD_ORIENT.LANDSCAPE
        _assert_mm(section.page_width, 297)
        _assert_mm(section.page_height, 210)
        assert (
            section.left_margin
            == section.right_margin
            == section.top_margin
            == section.bottom_margin
        )

    def test_configures_every_section(self):
        doc = Document()
        doc.add_section()
        configure_document(doc)
        for section in doc.sections:
            assert section.orientation == WD_ORIENT.LANDSCAPE
            _assert_mm(section.page_width, 297)
            _assert_mm(section.page_height, 210)


NAVY = RGBColor.from_string("1F3864")
TEAL = RGBColor.from_string("17635F")


class TestConfigureDocumentStyles:
    def test_normal_style_uses_calibri(self):
        doc = Document()
        configure_document(doc)
        normal = doc.styles["Normal"]
        assert normal.font.name == "Calibri"
        assert normal.font.size == Pt(10)

    def test_heading_hierarchy_colors_and_keep_with_next(self):
        doc = Document()
        configure_document(doc)
        h1 = doc.styles["Heading 1"]
        h2 = doc.styles["Heading 2"]
        assert h1.font.color.rgb == NAVY
        assert h1.paragraph_format.keep_with_next is True
        assert h2.font.color.rgb == TEAL
        assert h2.paragraph_format.keep_with_next is True

    def test_title_style_is_large_and_navy(self):
        doc = Document()
        configure_document(doc)
        title = doc.styles["Title"]
        assert title.font.color.rgb == NAVY
        assert title.font.size >= Pt(24)


class TestConfigureDocumentFurniture:
    def test_footer_has_page_and_numpages_fields(self):
        doc = Document()
        configure_document(doc)
        footer = doc.sections[0].footer
        xml = footer.paragraphs[0]._p.xml
        assert "PAGE" in xml
        assert "NUMPAGES" in xml
        assert "Página" in footer.paragraphs[0].text

    def test_header_shows_project_name_when_available(self):
        doc = Document()
        configure_document(doc, {"project": "Mina Norte"})
        header = doc.sections[0].header
        assert "Mina Norte" in header.paragraphs[0].text

    def test_header_falls_back_to_report_title(self):
        doc = Document()
        configure_document(doc)
        header = doc.sections[0].header
        assert "Informe de Conciliación Geotécnica" in header.paragraphs[0].text

    def test_first_page_has_no_running_furniture(self):
        doc = Document()
        configure_document(doc)
        section = doc.sections[0]
        assert section.different_first_page_header_footer is True
        assert section.first_page_header.is_linked_to_previous is True
        assert section.first_page_footer.is_linked_to_previous is True


class TestAddCover:
    def test_cover_shows_title_and_existing_metadata_labels(self):
        doc = Document()
        add_cover(doc, {"project": "Mina Norte", "author": "Geo Eng"})
        joined = "\n".join(p.text for p in doc.paragraphs)
        assert "Informe de Conciliación Geotécnica" in joined
        assert "Proyecto: Mina Norte" in joined
        assert "Elaborado por: Geo Eng" in joined
        assert f"Fecha: {datetime.now().strftime('%d/%m/%Y')}" in joined

    def test_cover_defaults_missing_metadata_to_na(self):
        doc = Document()
        add_cover(doc)
        joined = "\n".join(p.text for p in doc.paragraphs)
        assert "Proyecto: N/A" in joined
        assert "Elaborado por: N/A" in joined

    def test_cover_does_not_invent_metadata(self):
        doc = Document()
        add_cover(doc, {"project": "Mina Norte", "author": "Geo Eng"})
        joined = "\n".join(p.text for p in doc.paragraphs)
        for invented in ("Confidencial", "Aprobado", "Revisado", "Logo", "Empresa"):
            assert invented not in joined
        labeled = [p.text for p in doc.paragraphs if p.text.startswith(("Proyecto", "Elaborado", "Fecha"))]
        assert len(labeled) == 3

    def test_cover_preserves_long_metadata(self):
        doc = Document()
        long_project = "P" * 120
        long_author = "A" * 120
        add_cover(doc, {"project": long_project, "author": long_author})
        joined = "\n".join(p.text for p in doc.paragraphs)
        assert long_project in joined
        assert long_author in joined

    def test_cover_ends_with_page_break(self):
        doc = Document()
        add_cover(doc)
        doc.add_paragraph("contenido posterior")
        paragraphs = doc.paragraphs
        break_indexes = [
            i
            for i, p in enumerate(paragraphs)
            if p._p.findall(qn("w:r") + "/" + qn("w:br") + "[@" + qn("w:type") + "='page']")
        ]
        assert break_indexes
        fecha_indexes = [i for i, p in enumerate(paragraphs) if p.text.startswith("Fecha")]
        assert fecha_indexes
        assert max(fecha_indexes) < max(break_indexes)


class TestStyleDataTableHeader:
    def test_header_row_navy_fill_white_bold_text(self):
        doc = Document()
        table = _make_data_table(doc)
        style_data_table(table)
        for cell in table.rows[0].cells:
            assert _cell_fill(cell) == "1F3864"
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    assert run.font.bold is True
                    assert run.font.color.rgb == RGBColor.from_string("FFFFFF")

    def test_header_row_repeats_across_pages(self):
        doc = Document()
        table = _make_data_table(doc)
        style_data_table(table)
        assert _row_has_flag(table.rows[0], "w:tblHeader")


ZEBRA_FILL = "EEF3F8"


def _apply_generator_status_fill(cell, fill_hex, font_hex):
    cell.text = "✓ 85"
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shd)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(font_hex)


class TestStyleDataTableBody:
    def test_alternate_data_rows_get_pale_fill(self):
        doc = Document()
        table = _make_data_table(doc)
        style_data_table(table)
        assert _cell_fill(table.rows[1].cells[0]) is None
        assert _cell_fill(table.rows[2].cells[0]) == ZEBRA_FILL
        assert _cell_fill(table.rows[3].cells[0]) is None

    def test_existing_status_fill_and_font_color_preserved(self):
        doc = Document()
        table = _make_data_table(doc)
        status_cell = table.rows[2].cells[2]
        _apply_generator_status_fill(status_cell, "C6EFCE", "006100")
        style_data_table(table)
        assert _cell_fill(status_cell) == "C6EFCE"
        for paragraph in status_cell.paragraphs:
            for run in paragraph.runs:
                assert run.font.color.rgb == RGBColor.from_string("006100")
                assert run.font.bold is True

    def test_data_rows_do_not_split_across_pages(self):
        doc = Document()
        table = _make_data_table(doc)
        style_data_table(table)
        for row in table.rows:
            assert _row_has_flag(row, "w:cantSplit")

    def test_modest_cell_padding_applied(self):
        doc = Document()
        table = _make_data_table(doc)
        style_data_table(table)
        tblPr = table._tbl.tblPr
        cell_mar = tblPr.find(qn("w:tblCellMar"))
        assert cell_mar is not None
        left = cell_mar.find(qn("w:left"))
        right = cell_mar.find(qn("w:right"))
        assert left is not None and right is not None
        assert left.get(qn("w:w")) == "80"
        assert right.get(qn("w:w")) == "80"

    def test_fits_usable_width_with_column_weights(self):
        doc = Document()
        configure_document(doc)
        section = doc.sections[0]
        usable = section.page_width - section.left_margin - section.right_margin
        table = _make_data_table(doc)
        style_data_table(table, column_weights=[2, 1, 1])
        total = sum(col.width for col in table.columns)
        assert abs(total - usable) < Mm(1)
        widths = [col.width for col in table.columns]
        assert abs(widths[0] - usable * 0.5) < Mm(1)
        assert abs(widths[1] - usable * 0.25) < Mm(1)

    def test_fits_usable_width_equally_without_weights(self):
        doc = Document()
        configure_document(doc)
        section = doc.sections[0]
        usable = section.page_width - section.left_margin - section.right_margin
        table = _make_data_table(doc)
        style_data_table(table)
        assert table.autofit is False
        for col in table.columns:
            assert abs(col.width - usable / 3) < Mm(1)

    def test_numeric_cells_right_aligned(self):
        doc = Document()
        table = _make_data_table(doc)
        style_data_table(table)
        assert table.rows[1].cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
        assert table.rows[1].cells[0].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
        assert table.rows[3].cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT

    def test_data_cell_runs_get_compact_size_but_keep_color(self):
        doc = Document()
        table = _make_data_table(doc)
        style_data_table(table)
        cell = table.rows[1].cells[0]
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                assert run.font.size == Pt(9)
                assert run.font.color.rgb is None

    def test_empty_table_is_safe_noop(self):
        doc = Document()
        table = doc.add_table(rows=0, cols=3)
        style_data_table(table)
        assert len(table.rows) == 0


class TestTableWidthTwips:
    def test_tblw_stored_in_twips_matches_grid_and_usable_width(self, tmp_path):
        doc = Document()
        configure_document(doc)
        section = doc.sections[0]
        usable = section.page_width - section.left_margin - section.right_margin
        table = _make_data_table(doc)
        style_data_table(table, column_weights=[2, 1, 1])

        out = tmp_path / "tblw.docx"
        doc.save(str(out))
        reloaded = Document(str(out))
        styled = reloaded.tables[0]
        tblW = styled._tbl.tblPr.find(qn("w:tblW"))
        assert tblW is not None
        assert tblW.get(qn("w:type")) == "dxa"
        stored = int(tblW.get(qn("w:w")))
        grid = styled._tbl.find(qn("w:tblGrid"))
        grid_total = sum(int(col.get(qn("w:w"))) for col in grid.findall(qn("w:gridCol")))
        assert abs(stored - grid_total) <= 2
        assert abs(stored - Emu(int(usable)).twips) <= 2


class TestHeaderFooterDistance:
    def test_distances_explicit_inside_margins_with_safe_gap(self):
        doc = Document()
        configure_document(doc)
        section = doc.sections[0]
        margin = section.top_margin
        for distance in (section.header_distance, section.footer_distance):
            assert distance > 0
            assert distance < margin
            assert margin - distance >= Mm(5)


class TestHeaderAbbreviation:
    def test_long_project_name_abbreviated_in_header_only(self):
        doc = Document()
        long_name = "P" * 300
        configure_document(doc, {"project": long_name})
        add_cover(doc, {"project": long_name})
        header_text = doc.sections[0].header.paragraphs[0].text
        assert len(header_text) <= 60
        assert header_text.startswith("P")
        assert header_text.endswith("…")
        joined = "\n".join(p.text for p in doc.paragraphs)
        assert long_name in joined

    def test_header_normalizes_embedded_newlines(self):
        doc = Document()
        configure_document(doc, {"project": "Mina\nNorte   Sector  2"})
        header_text = doc.sections[0].header.paragraphs[0].text
        assert "\n" not in header_text
        assert "Mina Norte Sector 2" in header_text


class TestLayoutPipelineSmoke:
    def test_configure_cover_and_table_survive_save_and_reload(self, tmp_path):
        doc = Document()
        configure_document(doc, {"project": "Mina Norte", "author": "Geo Eng"})
        add_cover(doc, {"project": "Mina Norte", "author": "Geo Eng"})

        doc.add_heading("1. Resumen Ejecutivo", level=1)
        table = _make_data_table(doc)
        _apply_generator_status_fill(table.rows[1].cells[2], "C6EFCE", "006100")
        style_data_table(table, column_weights=[2, 1, 1])

        out = tmp_path / "layout_smoke.docx"
        doc.save(str(out))
        assert out.stat().st_size > 0

        reloaded = Document(str(out))
        joined = "\n".join(p.text for p in reloaded.paragraphs)
        assert "Informe de Conciliación Geotécnica" in joined
        assert "Proyecto: Mina Norte" in joined
        assert "Mina Norte" in reloaded.sections[0].header.paragraphs[0].text
        footer_xml = reloaded.sections[0].footer.paragraphs[0]._p.xml
        assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml
        styled = reloaded.tables[0]
        assert _cell_fill(styled.rows[0].cells[0]) == "1F3864"
        assert _cell_fill(styled.rows[1].cells[2]) == "C6EFCE"
        assert _row_has_flag(styled.rows[0], "w:tblHeader")
        assert _row_has_flag(styled.rows[2], "w:cantSplit")
