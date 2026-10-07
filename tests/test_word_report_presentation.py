"""Presentation and preservation tests for the generate_word_report layout.

Presentation tests (cover, furniture, table styling, two-column profile
grid, dual-bounded images) are expected to fail against the un-integrated
generator; preservation tests (exact cell content, status fills, N/A
fallbacks, comparison immutability) must pass before and after integration.
"""
import copy
import io
from datetime import datetime
from types import SimpleNamespace

import matplotlib
matplotlib.use("Agg")
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.image.image import Image as DocxImage
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Mm
from PIL import Image as PILImage

from core.calculo_tronadura import procesar_pozos
from core.param_extractor import BenchParams, ExtractionResult, compare_design_vs_asbuilt
from core.report_generator import (
    _add_bounded_picture,
    create_plan_view_image,
    create_section_plot,
    generate_word_report,
)
from core.section_cutter import SectionLine

TOLERANCES = {
    "bench_height": {"neg": 1.0, "pos": 1.5},
    "face_angle": {"neg": 5.0, "pos": 5.0},
    "berm_width": {"min": 6.0},
}

GREEN_FILL = "C6EFCE"
RED_FILL = "FFC7CE"
NAVY_FILL = "1F3864"
MM_TOLERANCE = Mm(0.5)
PROFILE_MAX_HEIGHT = Inches(2.81)
PROFILE_MAX_WIDTH = Inches(5.15)
PLAN_MAX_HEIGHT = Inches(5.5)
PLAN_MAX_WIDTH = Inches(10.45)


def _bench(num, crest_z, toe_z, crest_x, toe_x, face=70.0, berm=9.0):
    return BenchParams(
        bench_number=num,
        crest_elevation=float(crest_z),
        crest_distance=float(crest_x),
        toe_elevation=float(toe_z),
        toe_distance=float(toe_x),
        bench_height=float(abs(crest_z - toe_z)),
        face_angle=float(face),
        berm_width=float(berm),
    )


def _extraction(section_name, benches):
    return ExtractionResult(section_name=section_name, sector="Test", benches=benches)


def _matched_comparisons(section_name="S-01"):
    design = _extraction(
        section_name,
        [_bench(1, 100, 85, 10, 20), _bench(2, 85, 70, 30, 40)],
    )
    topo = _extraction(
        section_name,
        [
            _bench(1, 99, 83, 11, 21, face=68.0, berm=8.5),
            _bench(2, 84, 69, 31, 41, face=72.0, berm=10.0),
        ],
    )
    return compare_design_vs_asbuilt(design, topo, TOLERANCES)


def _all_data_item(section_name="S-01"):
    d = np.linspace(0, 50, 25)
    return {
        "section_name": section_name,
        "params_design": None,
        "params_topo": None,
        "profile_d": (d, 100.0 - d * 0.5),
        "profile_t": (d, 99.0 - d * 0.5),
    }


def _section_line(name):
    return SectionLine(name=name, origin=np.array([0.0, 0.0]), azimuth=90.0, length=200.0)


def _known_comparisons():
    return [
        {
            "type": "MATCH",
            "section": "S-01",
            "bench_num": 1,
            "level": "85",
            "height_status": "CUMPLE",
            "angle_status": "CUMPLE",
            "berm_status": "CUMPLE",
            "height_design": 15.0,
            "height_real": 14.2,
            "angle_design": 70.0,
            "angle_real": 69.0,
            "berm_design": 9.0,
            "berm_real": 8.5,
            "bench_score": 85.0,
            "bench_real": SimpleNamespace(floor_elevation=80.0, crest_elevation=100.0),
        },
        {
            "type": "MATCH",
            "section": "S-01",
            "bench_num": 2,
            "level": "70",
            "height_status": "NO CUMPLE",
            "angle_status": "NO CUMPLE",
            "berm_status": "NO CUMPLE",
            "height_design": 15.0,
            "height_real": 16.8,
            "angle_design": 70.0,
            "angle_real": 61.0,
            "berm_design": 9.0,
            "berm_real": 4.0,
            "bench_score": 40.0,
            "bench_real": SimpleNamespace(floor_elevation=70.0, crest_elevation=85.0),
        },
    ]


def _blast_holes_df():
    raw = pd.DataFrame(
        [
            {
                "Latitud_Geo": 10.0,
                "Longitud_Geo": 0.0,
                "Nombre_Banco": 4000.0,
                "Inclinacion_real": 0.0,
                "Azimuth_real": 0.0,
                "longitud_real": 10.0,
                "Kilos_Cargados_real": 200.0,
            }
        ]
    )
    return procesar_pozos(raw, incl_convention="from_vertical", geometry_user_confirmed=True)[0]


def _cell_fill(cell):
    tcPr = cell._tc.tcPr
    if tcPr is None:
        return None
    for child in tcPr:
        if child.tag == qn("w:shd"):
            return child.get(qn("w:fill"))
    return None


def _row_has_header_flag(row):
    trPr = row._tr.trPr
    return trPr is not None and trPr.find(qn("w:tblHeader")) is not None


def _table_by_header_cell(doc, header_text):
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text == header_text:
            return table
    return None


def _image_tables(doc):
    tables = []
    for table in doc.tables:
        has_image = any(
            cell._tc.find(".//" + qn("w:drawing")) is not None
            for row in table.rows
            for cell in row.cells
        )
        if has_image:
            tables.append(table)
    return tables


def _extents(element):
    return [
        (int(ext.get("cx")), int(ext.get("cy")))
        for ext in element.iter(qn("wp:extent"))
    ]


def _has_page_break(paragraph):
    return bool(
        paragraph._p.findall(
            qn("w:r") + "/" + qn("w:br") + "[@" + qn("w:type") + "='page']"
        )
    )


def _native_size(png_bytes):
    image = DocxImage.from_blob(png_bytes)
    dpi_w = image.horz_dpi if image.horz_dpi and image.horz_dpi > 0 else 96.0
    dpi_h = image.vert_dpi if image.vert_dpi and image.vert_dpi > 0 else 96.0
    return image.px_width / dpi_w, image.px_height / dpi_h


def _generate(tmp_path, comparisons, all_data, **kwargs):
    out = tmp_path / "report.docx"
    generate_word_report(comparisons, all_data, str(out), **kwargs)
    return Document(str(out))


class TestCoverAndFurniture:
    def test_landscape_a4_with_uniform_margins(self, tmp_path):
        doc = _generate(tmp_path, [], [])
        section = doc.sections[0]
        assert section.orientation == WD_ORIENT.LANDSCAPE
        assert abs(section.page_width - Mm(297)) < MM_TOLERANCE
        assert abs(section.page_height - Mm(210)) < MM_TOLERANCE
        assert (
            section.left_margin
            == section.right_margin
            == section.top_margin
            == section.bottom_margin
        )

    def test_cover_keeps_exact_title_and_metadata_paragraphs(self, tmp_path):
        doc = _generate(
            tmp_path,
            [],
            [],
            project_info={"project": "Mina Norte", "author": "Geo Eng"},
        )
        texts = [p.text for p in doc.paragraphs]
        assert "Informe de Conciliación Geotécnica" in texts
        assert "Proyecto: Mina Norte" in texts
        assert "Elaborado por: Geo Eng" in texts
        assert f"Fecha: {datetime.now().strftime('%d/%m/%Y')}" in texts

    def test_cover_na_fallback_and_blank_strings_preserved(self, tmp_path):
        doc = _generate(tmp_path, [], [])
        texts = [p.text for p in doc.paragraphs]
        assert "Proyecto: N/A" in texts
        assert "Elaborado por: N/A" in texts

        doc_blank = _generate(
            tmp_path, [], [], project_info={"project": "", "author": ""}
        )
        blank_texts = [p.text for p in doc_blank.paragraphs]
        assert "Proyecto: " in blank_texts
        assert "Elaborado por: " in blank_texts

    def test_cover_preserves_long_metadata(self, tmp_path):
        long_project = "P" * 120
        long_author = "A" * 120
        doc = _generate(
            tmp_path,
            [],
            [],
            project_info={"project": long_project, "author": long_author},
        )
        texts = [p.text for p in doc.paragraphs]
        assert f"Proyecto: {long_project}" in texts
        assert f"Elaborado por: {long_author}" in texts

    def test_footer_has_page_fields_and_header_shows_project(self, tmp_path):
        doc = _generate(
            tmp_path,
            [],
            [],
            project_info={"project": "Mina Norte"},
        )
        section = doc.sections[0]
        footer_xml = section.footer.paragraphs[0]._p.xml
        assert "PAGE" in footer_xml
        assert "NUMPAGES" in footer_xml
        assert "Mina Norte" in section.header.paragraphs[0].text

    def test_first_page_has_no_running_furniture(self, tmp_path):
        doc = _generate(tmp_path, [], [])
        section = doc.sections[0]
        assert section.different_first_page_header_footer is True
        assert section.first_page_header.is_linked_to_previous is True
        assert section.first_page_footer.is_linked_to_previous is True


class TestDataTables:
    def test_summary_table_header_repeats_with_navy_fill(self, tmp_path):
        doc = _generate(tmp_path, _known_comparisons(), [])
        table = _table_by_header_cell(doc, "Sección")
        assert table is not None
        assert _row_has_header_flag(table.rows[0])
        for cell in table.rows[0].cells:
            assert _cell_fill(cell) == NAVY_FILL

    def test_parameter_and_depth_tables_have_styled_headers(self, tmp_path):
        doc = _generate(tmp_path, _known_comparisons(), [])
        for header_text in ("Parámetro", "Métrica"):
            table = _table_by_header_cell(doc, header_text)
            assert table is not None, header_text
            assert _row_has_header_flag(table.rows[0])
            for cell in table.rows[0].cells:
                assert _cell_fill(cell) == NAVY_FILL

    def test_blast_correlation_table_header_styled(self, tmp_path):
        comps = _known_comparisons()
        doc = _generate(
            tmp_path,
            comps,
            [],
            df_pozos=_blast_holes_df(),
            sections=[_section_line("S-01")],
        )
        correlation = None
        for table in doc.tables:
            if table.rows and table.rows[0].cells[1].text == "Pozos Cercanos":
                correlation = table
                break
        assert correlation is not None
        assert _row_has_header_flag(correlation.rows[0])
        for cell in correlation.rows[0].cells:
            assert _cell_fill(cell) == NAVY_FILL

    def test_profile_grid_never_gets_data_table_styling(self, tmp_path):
        comparisons, all_data, sections = [], [], []
        for i in range(2):
            name = f"S-{i:02d}"
            comparisons.extend(_matched_comparisons(name))
            all_data.append(_all_data_item(name))
            sections.append(_section_line(name))
        doc = _generate(tmp_path, comparisons, all_data, sections=sections)
        for table in _image_tables(doc):
            assert not any(_cell_fill(cell) == NAVY_FILL for row in table.rows for cell in row.cells)

    def test_status_text_and_fills_preserved(self, tmp_path):
        doc = _generate(tmp_path, _known_comparisons(), [])
        table = _table_by_header_cell(doc, "Sección")
        expected = [("✓ 85", GREEN_FILL, "006100"), ("✗ 40", RED_FILL, "9C0006")]
        for row, (text, fill, font_rgb) in zip(table.rows[1:], expected):
            estado = row.cells[10]
            assert estado.text == text
            assert _cell_fill(estado) == fill
            run = estado.paragraphs[0].runs[0]
            assert run.font.color.rgb is not None
            assert str(run.font.color.rgb) == font_rgb

    def test_summary_cell_content_exact(self, tmp_path):
        doc = _generate(tmp_path, _known_comparisons(), [])
        table = _table_by_header_cell(doc, "Sección")
        assert [c.text for c in table.rows[1].cells] == [
            "S-01", "B1 (85)", "15.0", "14.2", "70.0", "69.0",
            "9.0", "8.5", "14.2", "80.0", "✓ 85",
        ]
        assert [c.text for c in table.rows[2].cells] == [
            "S-01", "B2 (70)", "15.0", "16.8", "70.0", "61.0",
            "9.0", "4.0", "16.8", "70.0", "✗ 40",
        ]

    def test_depth_and_parameter_cell_content_exact(self, tmp_path):
        doc = _generate(tmp_path, _known_comparisons(), [])
        depth = _table_by_header_cell(doc, "Métrica")
        assert [[c.text for c in row.cells] for row in depth.rows] == [
            ["Métrica", "Valor"],
            ["Cota Piso Global", "70.00 m"],
            ["Cota Cresta Global", "100.00 m"],
            ["Profundidad Total", "30.00 m"],
        ]
        params = _table_by_header_cell(doc, "Parámetro")
        assert [[c.text for c in row.cells] for row in params.rows] == [
            ["Parámetro", "CUMPLE", "NO CUMPLE", "% Logro", "Valor Promedio (Real)"],
            ["Altura", "1", "1", "50.0%", "15.50 m"],
            ["Ángulo Cara", "1", "1", "50.0%", "65.00 °"],
            ["Berma", "1", "1", "50.0%", "6.25 m"],
        ]

    def test_status_fills_continue_on_every_row(self, tmp_path):
        comparisons, all_data = [], []
        for i in range(12):
            name = f"S-{i:02d}"
            comparisons.extend(_matched_comparisons(name))
            all_data.append(_all_data_item(name))
        doc = _generate(tmp_path, comparisons, all_data)
        table = _table_by_header_cell(doc, "Sección")
        data_rows = table.rows[1:]
        assert len(data_rows) >= 9
        for row in data_rows:
            estado = row.cells[10]
            cumple = estado.text.startswith("✓")
            expected_fill = GREEN_FILL if cumple else RED_FILL
            assert _cell_fill(estado) == expected_fill, estado.text

    def test_comparison_inputs_unchanged(self, tmp_path):
        comps = _known_comparisons()
        snapshot = copy.deepcopy([{k: v for k, v in c.items()} for c in comps])
        _generate(tmp_path, comps, [_all_data_item("S-01")])
        assert comps == snapshot


class TestProfileGrid:
    def _multi_section_fixture(self, n_sections):
        comparisons, all_data = [], []
        for i in range(n_sections):
            name = f"S-{i:02d}"
            comparisons.extend(_matched_comparisons(name))
            all_data.append(_all_data_item(name))
        return comparisons, all_data

    def test_two_columns_four_profiles_per_page(self, tmp_path):
        comparisons, all_data = self._multi_section_fixture(9)
        doc = _generate(tmp_path, comparisons, all_data)
        image_tables = _image_tables(doc)
        assert len(image_tables) == 3
        total_images = 0
        for table in image_tables:
            assert len(table.columns) == 2
            assert len(table.rows) <= 2
            total_images += sum(len(_extents(cell._tc)) for row in table.rows for cell in row.cells)
        assert total_images == 9

    def test_profile_images_dual_bounded(self, tmp_path):
        comparisons, all_data = self._multi_section_fixture(4)
        doc = _generate(tmp_path, comparisons, all_data)
        for table in _image_tables(doc):
            for row in table.rows:
                for cell in row.cells:
                    for cx, cy in _extents(cell._tc):
                        assert cx <= PROFILE_MAX_WIDTH
                        assert cy <= PROFILE_MAX_HEIGHT
                        assert cy >= Inches(2.7) or cx >= Inches(4.9)

    def test_profile_aspect_never_distorted(self, tmp_path):
        comparisons = _matched_comparisons("S-01")
        all_data = [_all_data_item("S-01")]
        doc = _generate(tmp_path, comparisons, all_data)
        item = all_data[0]
        reference = create_section_plot(
            None,
            None,
            item["profile_d"][0],
            item["profile_d"][1],
            item["profile_t"][0],
            item["profile_t"][1],
            plot_options={},
            section=None,
            df_pozos=None,
            filtered_bench_nums={c["bench_num"] for c in comparisons},
        )
        native_w, native_h = _native_size(reference.getvalue())
        reference.close()
        for table in _image_tables(doc):
            for row in table.rows:
                for cell in row.cells:
                    for cx, cy in _extents(cell._tc):
                        assert abs((cx / cy) - (native_w / native_h)) / (native_w / native_h) < 0.01

    def test_tall_and_wide_profiles_stay_bounded(self, tmp_path):
        d = np.linspace(0, 400, 200)
        wide_item = {
            "section_name": "S-WIDE",
            "params_design": None,
            "params_topo": None,
            "profile_d": (d, 100.0 - d * 0.05),
            "profile_t": (d, 99.0 - d * 0.05),
        }
        comparisons = _matched_comparisons("S-WIDE")
        doc = _generate(tmp_path, comparisons, [wide_item])
        for table in _image_tables(doc):
            for row in table.rows:
                for cell in row.cells:
                    for cx, cy in _extents(cell._tc):
                        assert cx <= PROFILE_MAX_WIDTH
                        assert cy <= PROFILE_MAX_HEIGHT


class TestBoundedPictureUnitConversion:
    def _png(self, px_w, px_h, dpi):
        buf = io.BytesIO()
        PILImage.new("RGB", (px_w, px_h), "white").save(buf, format="PNG", dpi=(dpi, dpi))
        buf.seek(0)
        return buf

    def _embed(self, px_w, px_h, dpi, max_w, max_h):
        doc = Document()
        paragraph = doc.add_paragraph()
        _add_bounded_picture(paragraph, self._png(px_w, px_h, dpi), max_w, max_h)
        extents = _extents(paragraph._p)
        assert len(extents) == 1
        return extents[0]

    def test_width_limited_image_uses_emu_conversion(self):
        cx, cy = self._embed(1600, 800, 200, Inches(3.0), Inches(2.0))
        assert abs(cx - Inches(3.0)) <= Emu(1000)
        assert abs(cy - Inches(1.5)) <= Emu(1000)
        assert cx <= Inches(3.0)
        assert cy <= Inches(2.0)
        assert abs((cx / cy) - 2.0) < 0.01

    def test_height_limited_image_uses_emu_conversion(self):
        cx, cy = self._embed(800, 1600, 100, Inches(3.0), Inches(4.0))
        assert abs(cy - Inches(4.0)) <= Emu(1000)
        assert abs(cx - Inches(2.0)) <= Emu(1000)
        assert cx <= Inches(3.0)
        assert cy <= Inches(4.0)
        assert abs((cx / cy) - 0.5) < 0.01


class TestPlanViewImage:
    def test_plan_image_bounded_on_dedicated_page(self, tmp_path):
        comparisons = _matched_comparisons("S-01")
        all_data = [_all_data_item("S-01")]
        sections = [_section_line("S-01")]
        doc = _generate(tmp_path, comparisons, all_data, sections=sections)
        plan_paragraphs = [
            i
            for i, p in enumerate(doc.paragraphs)
            if _extents(p._p) and i > 0 and _has_page_break(doc.paragraphs[i - 1])
        ]
        assert len(plan_paragraphs) == 1
        idx = plan_paragraphs[0]
        plan_extents = _extents(doc.paragraphs[idx]._p)
        assert plan_extents
        for cx, cy in plan_extents:
            assert cx <= PLAN_MAX_WIDTH
            assert cy <= PLAN_MAX_HEIGHT
        assert _has_page_break(doc.paragraphs[idx + 1])
        reference = create_plan_view_image(comparisons, sections)
        native_w, native_h = _native_size(reference.getvalue())
        reference.close()
        for cx, cy in plan_extents:
            assert abs((cx / cy) - (native_w / native_h)) / (native_w / native_h) < 0.01


class TestEmptyAndEdgeInputs:
    def test_empty_inputs_still_produce_cover_and_valid_document(self, tmp_path):
        doc = _generate(tmp_path, [], [])
        texts = [p.text for p in doc.paragraphs]
        assert "Informe de Conciliación Geotécnica" in texts
        assert "Proyecto: N/A" in texts
        assert "No se encontraron resultados para reportar." in texts
        assert "No hay datos de comparación disponibles." in texts

    def test_document_without_sections_has_no_plan_image(self, tmp_path):
        comparisons = _matched_comparisons("S-01")
        doc = _generate(tmp_path, comparisons, [_all_data_item("S-01")])
        assert not [cy for _, cy in _extents(doc.element.body) if cy >= Inches(3)]
